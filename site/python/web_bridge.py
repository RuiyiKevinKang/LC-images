"""Bounded browser/file adapter for the unchanged, published LC solver.

The public functions return ordinary JSON-serializable Python values. Numerical
binning, feature construction, time alignment and solving belong to the two
original source files; this module only validates, adapts and exports data.
"""
import csv
import hashlib
import json
import math
import time
import zipfile
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from bin_pixels import (build_features, finite_array, prepare_green_images,
                        prepare_lc, validate_time, geometry_from_xy, _video_ids)
from solve_A import (prepare_target, input_pseudoinverse, solve_mapping,
                     check_requirements)

MAX_FILE_BYTES = 40 * 1024 * 1024
MAX_ELEMENTS = 2_000_000
MAX_FRAMES = 600
MAX_XLSX_CELLS = 500_000
MAX_XLSX_EXPORT = 200_000
_MIMES = {".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
          ".csv": "text/csv;charset=utf-8", ".json": "application/json",
          ".npy": "application/octet-stream", ".npz": "application/octet-stream",
          ".zip": "application/zip"}


def _plain(value):
    if isinstance(value, np.ndarray):
        return _plain(value.tolist())
    if isinstance(value, np.generic):
        return _plain(value.item())
    if isinstance(value, dict):
        return {str(k): _plain(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_plain(v) for v in value]
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def _json(path, value):
    Path(path).write_text(json.dumps(_plain(value), ensure_ascii=False, indent=2,
                                    allow_nan=False), encoding="utf-8")


def _download(path):
    path = Path(path)
    return {"name": path.name, "path": str(path),
            "mime": _MIMES.get(path.suffix, "application/octet-stream")}


def _real(value, name):
    original = np.asarray(value)
    if original.dtype.kind in "cObSU":
        # JSON numbers and numeric table cells should be genuine real numbers.
        # Object arrays and strings from scientific archives are never executed.
        raise ValueError(f"{name} 必须由实数构成，不能使用文本、布尔值、复数或对象数组")
    if original.size > MAX_ELEMENTS:
        raise ValueError(f"{name} 超过 {MAX_ELEMENTS:,} 个数值的浏览器计算上限")
    return finite_array(original, name)


def _number(value, name):
    if isinstance(value, bool):
        raise ValueError(f"{name} 必须为有限实数")
    try:
        result = float(value)
    except (TypeError, ValueError, OverflowError):
        raise ValueError(f"{name} 必须为有限实数") from None
    if not math.isfinite(result):
        raise ValueError(f"{name} 必须为有限实数")
    return result


def _check_zip(path, limit=64 * 1024 * 1024):
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 200 or sum(item.file_size for item in entries) > limit:
            raise ValueError("压缩文件解压后的体积或文件数量超过浏览器限制")


def _check_npy_header(stream):
    version = np.lib.format.read_magic(stream)
    if version == (1, 0):
        shape, _, dtype = np.lib.format.read_array_header_1_0(stream)
    elif version in ((2, 0), (3, 0)):
        shape, _, dtype = np.lib.format.read_array_header_2_0(stream)
    else:
        raise ValueError("无法识别 NPY 文件版本")
    if dtype.hasobject:
        raise ValueError("不接受需要 pickle 的对象数组")
    if len(shape) > 4 or math.prod(shape) > MAX_ELEMENTS:
        raise ValueError("数组维数或元素数量超过浏览器限制（最多 2,000,000 个元素）")


def _metadata(rows):
    result = {}
    for row in rows:
        if not row or row[0] is None or str(row[0]).strip().lower() in ("key", "参数"):
            continue
        if len(row) < 2:
            continue
        key, value = str(row[0]).strip(), row[1]
        if key not in ("time", "LC_time_s", "model_time", "kind", "cell_area",
                       "image_shape", "video_ids", "time_is_frame_index"):
            continue
        if isinstance(value, str):
            try:
                value = json.loads(value)
            except json.JSONDecodeError:
                pass
        result[key] = value
    return result


def _read_rows(path):
    if path.suffix.lower() == ".xlsx":
        _check_zip(path, 64 * 1024 * 1024)
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheets = [sheet for sheet in workbook.worksheets
                      if sheet.title.lower() not in ("metadata", "readme", "说明")]
            if not sheets:
                raise ValueError("工作簿中没有数据工作表")
            sheet = sheets[0]
            if (sheet.max_row or 1) * (sheet.max_column or 1) > MAX_XLSX_CELLS:
                raise ValueError("XLSX 数据表超过 500,000 个单元格；大数组请使用 NPY/NPZ")
            rows = []
            count = 0
            for row in sheet.iter_rows(values_only=True):
                count += len(row)
                if count > MAX_XLSX_CELLS:
                    raise ValueError("XLSX 数据表超过 500,000 个单元格")
                if any(value is not None for value in row):
                    rows.append(list(row))
            metadata = {}
            for extra in workbook.worksheets:
                if extra.title.lower() == "metadata":
                    if (extra.max_row or 1) * (extra.max_column or 1) > 10_000:
                        raise ValueError("metadata 工作表过大")
                    metadata = _metadata(extra.iter_rows(values_only=True))
            return rows, metadata
        finally:
            workbook.close()
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        rows = []
        cells = 0
        for row in csv.reader(stream, delimiter=delimiter):
            cells += len(row)
            if cells > MAX_ELEMENTS * 5:
                raise ValueError("表格单元格数量超过浏览器限制")
            if any(value.strip() for value in row):
                rows.append(row)
    return rows, {}


def _table_source(rows, role):
    if len(rows) < 2:
        raise ValueError("数据表必须包含列名和至少一行数值")
    headers = [str(v).strip() if v is not None else "" for v in rows[0]]
    while headers and not headers[-1]:
        headers.pop()
    if len(set(headers)) != len(headers) or not all(headers):
        raise ValueError("列名不能为空或重复")
    data = []
    for index, row in enumerate(rows[1:], 2):
        if len(row) > len(headers) and any(v not in (None, "") for v in row[len(headers):]):
            raise ValueError(f"第 {index} 行的数据列数不一致")
        row = list(row[:len(headers)])
        if len(row) != len(headers) or any(v is None or v == "" for v in row):
            raise ValueError(f"第 {index} 行有缺失值；不允许自动补零")
        data.append([_number(v, f"第 {index} 行") for v in row])
    table = np.asarray(data, dtype=float)
    lookup = {h.lower(): i for i, h in enumerate(headers)}
    key = role.lower()
    time_name = next((name for name in ("t", "time", "time_s", "lc_time_s", "frame")
                      if name in lookup), None)
    if all(name in lookup for name in ("x", "y", key)) and time_name:
        ti, xi, yi, vi = [lookup[n] for n in (time_name, "x", "y", key)]
        times = np.unique(table[:, ti])
        if len(times) > MAX_FRAMES:
            raise ValueError("最多支持 600 帧")
        coordinates = sorted(set(map(tuple, table[:, [xi, yi]])))
        if len(times) * len(coordinates) != len(table):
            raise ValueError("长表中每一帧必须包含相同坐标，且 (t,x,y) 不能重复")
        coord_index = {coord: i for i, coord in enumerate(coordinates)}
        time_index = {t: i for i, t in enumerate(times)}
        values = np.empty((len(times), len(coordinates)))
        seen = set()
        for row in table:
            cell = (time_index[row[ti]], coord_index[(row[xi], row[yi])])
            if cell in seen:
                raise ValueError("长表含重复 (t,x,y)；不能平均或覆盖重复数据")
            seen.add(cell)
            values[cell] = row[vi]
        if len(seen) != values.size:
            raise ValueError("长表中每一帧必须包含完整的固定坐标集合")
        source = {role: values, "xy": np.asarray(coordinates, dtype=float)}
        if time_name != "frame":
            source["time"] = times
        return source
    if role == "G" and headers[:3] == ["pixel_id", "x", "y"]:
        frames = len(headers) - 3
        if frames < 1 or headers[3:] != [f"G_frame_{t:03d}" for t in range(frames)]:
            raise ValueError("G 宽表列名须为 pixel_id,x,y,G_frame_000,G_frame_001,...")
        ids = table[:, 0]
        if not np.all(ids == np.floor(ids)) or len(np.unique(ids)) != len(ids):
            raise ValueError("pixel_id 必须为不重复的整数")
        return {"G": table[:, 3:].T, "xy": table[:, 1:3], "pixel_ids": ids}
    if role == "P":
        start = 1 if time_name and lookup[time_name] == 0 else 0
        expected = [f"p{i}" for i in range(len(headers) - start)]
        if [h.lower() for h in headers[start:]] == expected and expected:
            source = {"P": table[:, start:]}
            if start and time_name != "frame":
                source["time"] = table[:, 0]
            return source
    raise ValueError("无法识别表头：使用 t,x,y,G / t,x,y,P 长表；G 像素宽表；或 time,p0,p1,... 概率宽表")


def _load(path, role, options):
    path = Path(path)
    if not path.is_file():
        raise ValueError("输入文件不存在")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("单个输入文件不能超过 40 MB")
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".csv", ".tsv"):
        rows, meta = _read_rows(path)
        source = {**meta, **_table_source(rows, role)}
    elif suffix == ".json":
        source = json.loads(path.read_text(encoding="utf-8-sig"))
        if not isinstance(source, dict):
            raise ValueError("JSON 顶层必须为包含 G 或 P 键的对象")
    elif suffix == ".npy":
        with path.open("rb") as stream:
            _check_npy_header(stream)
        source = {role: np.load(path, allow_pickle=False)}
    elif suffix == ".npz":
        _check_zip(path)
        with zipfile.ZipFile(path) as archive:
            for member in archive.namelist():
                if member.endswith(".npy"):
                    with archive.open(member) as stream:
                        _check_npy_header(stream)
        with np.load(path, allow_pickle=False) as archive:
            source = {name: archive[name] for name in archive.files}
    else:
        raise ValueError("支持的输入格式为 XLSX、CSV、TSV、JSON、NPY 和 NPZ")
    for key in (role, "mean_green", "xy", "time", "LC_time_s", "model_time",
                "cell_area", "target_U", "target_L0", "mask"):
        if key in source:
            value = source[key]
            if key == "mask" and np.asarray(value).dtype.kind == "b":
                value = np.asarray(value, dtype=np.uint8)
            source[key] = _real(value, key)
    if role not in source and not (role == "G" and "mean_green" in source):
        raise ValueError(f"输入文件缺少 {role} 数组")
    if role == "G" and "G" in source and "xy" in source:
        if source["G"].ndim not in (2, 3):
            raise ValueError("带 xy 的 G 必须为 (T,N) 或 (V,T,N)")
    if "xy" in source:
        xy = source["xy"]
        if xy.ndim != 2 or xy.shape[1] != 2 or not len(xy):
            raise ValueError("xy 必须为非空的 (N,2) 坐标数组")
        if len(np.unique(xy, axis=0)) != len(xy):
            raise ValueError("xy 含重复坐标")
    if role == "P":
        if "target_kind" in options and options["target_kind"]:
            source["kind"] = options["target_kind"]
        if "cell_area" in options and options["cell_area"] not in (None, ""):
            area = options["cell_area"]
            if isinstance(area, str):
                try:
                    area = json.loads(area)
                except json.JSONDecodeError:
                    raise ValueError("cell_area 须为正数或 JSON 面积数组") from None
            source["cell_area"] = _real(area, "cell_area")
    return source


def _prepare_g(source, options):
    width = _number(options.get("bin_width", 16), "bin_width")
    if width <= 0:
        raise ValueError("bin_width 必须为正数")
    if "mean_green" in source:
        data = source["mean_green"]
        if data.ndim == 2:
            data = data[None]
        if data.ndim != 3 or min(data.shape) < 1:
            raise ValueError("mean_green 必须为 (T,B) 或 (V,T,B)")
        frames = data.shape[1]
        if frames > MAX_FRAMES or data.size * 3 > MAX_ELEMENTS:
            raise ValueError("最多支持 600 帧且三帧历史特征不能超过 2,000,000 个数值")
        index_time = "time" not in source
        lc = {"mean_green": data, "F": build_features(data),
              "time": validate_time(source.get("time", np.arange(frames)), frames),
              "time_is_frame_index": index_time, "geometry": None,
              "video_ids": _video_ids(source.get("video_ids", [f"v{i+1:03d}" for i in range(len(data))]), len(data))}
    else:
        G = source["G"]
        if "xy" in source:
            frames = G.shape[-2]
        else:
            if G.ndim not in (3, 4):
                raise ValueError("不含 xy 的 G 图像必须为 (T,H,W) 或 (V,T,H,W)")
            frames = G.shape[-3]
        if frames > MAX_FRAMES:
            raise ValueError("最多支持 600 帧")
        if "xy" in source:
            bound_xy = source["xy"]
        else:
            height, image_width = G.shape[-2:]
            mask = source.get("mask", np.ones((height, image_width)))
            if (mask.shape != (height, image_width) or not np.isin(mask, [0, 1]).all()
                    or not mask.any()):
                raise ValueError("ROI mask 必须为与 G 图像同尺寸的非空固定 0/1 掩膜")
            yy, xx = np.nonzero(mask)
            bound_xy = np.column_stack([xx, yy])
        bins = len(geometry_from_xy(bound_xy, width)["keys"])
        videos = G.shape[0] if G.ndim == (3 if "xy" in source else 4) else 1
        if videos * frames * bins * 3 > MAX_ELEMENTS:
            raise ValueError("三帧历史特征超过 2,000,000 个数值；请减少分箱数量或视频数量")
        if "xy" in source:
            lc = prepare_lc(G, source["xy"], source.get("time", np.arange(frames)),
                            width, source.get("video_ids"))
            lc["time_is_frame_index"] = "time" not in source
            if "pixel_ids" in source:
                lc["geometry"]["pixel_ids"] = source["pixel_ids"].astype(np.int64)
        else:
            lc = prepare_green_images(G, source.get("mask"), source.get("time"),
                                      width, source.get("video_ids"))
    if source.get("time_is_frame_index", False):
        lc["time_is_frame_index"] = True
    video_value = options.get("video_index", 1)
    video = _number(video_value, "video_index")
    if video != int(video) or video < 1 or video > len(lc["F"]):
        raise ValueError(f"video_index 必须为 1..{len(lc['F'])} 的整数")
    lc["selected_video_index"] = int(video)
    return lc


def inspect_input(path, role="G", options=None):
    """Validate one upload and expose dimensions; validation failures raise ValueError."""
    options = options or {}
    role = str(role).upper()
    if role not in ("G", "P"):
        raise ValueError("role 必须为 G 或 P")
    source = _load(path, role, options)
    if role == "G":
        lc = _prepare_g(source, options)
        array = source.get("G", source.get("mean_green"))
        geometry = lc["geometry"]
        result = {"frames": lc["F"].shape[1], "bins": lc["mean_green"].shape[2],
                  "features": lc["F"].shape[2], "videos": len(lc["F"]),
                  "cells": len(geometry["original_xy"]) if geometry else None,
                  "video_index": lc["selected_video_index"],
                  "time_is_frame_index": lc["time_is_frame_index"],
                  "is_binned_cache": geometry is None}
    else:
        array = source["P"]
        if array.ndim not in (2, 3) or min(array.shape) < 1:
            raise ValueError("P 必须为 (T,K) 或 (T,H,W)")
        if len(array) > MAX_FRAMES:
            raise ValueError("最多支持 600 帧")
        if np.any(array < 0):
            raise ValueError("P 含负值")
        if "xy" in source and len(source["xy"]) != int(np.prod(array.shape[1:])):
            raise ValueError("P 的 xy 坐标数量与网格数不一致")
        tv = source.get("LC_time_s", source.get("time", source.get("model_time")))
        if tv is not None:
            validate_time(tv, len(array), "P time")
        result = {"frames": len(array), "cells": int(np.prod(array.shape[1:])),
                  "videos": 1, "time_is_frame_index": tv is None,
                  "source_kind": str(np.asarray(source.get("kind", "mass")).item())}
        if result["source_kind"] == "density" and "cell_area" not in source:
            raise ValueError("密度输入必须明确提供 cell_area")
    return _plain({"status": "ready", "role": role, "filename": Path(path).name,
                   "shape": array.shape, "format": Path(path).suffix[1:].lower(),
                   "warnings": [], **result})


def _write_csv(path, header, rows):
    with Path(path).open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(header)
        writer.writerows(rows)


def _features(lc):
    bins = lc["mean_green"].shape[2]
    geometry = lc["geometry"]
    records = []
    for lag in (0, 1, 2):
        for block in range(bins):
            record = {"feature_index": lag * bins + block, "history_lag": lag,
                      "bin_index": block, "kx": None, "ky": None,
                      "pixel_count": None}
            if geometry:
                record.update(kx=int(geometry["keys"][block, 0]),
                              ky=int(geometry["keys"][block, 1]),
                              pixel_count=int(geometry["counts"][block]))
            records.append(record)
    return records


def _target_xy(source):
    if "xy" in source:
        return source["xy"]
    if source["P"].ndim == 3:
        yy, xx = np.indices(source["P"].shape[1:])
        return np.column_stack([xx.ravel(), yy.ravel()])
    return np.column_stack([np.arange(source["P"].shape[1]), np.zeros(source["P"].shape[1])])


def _source_info(path):
    path = Path(path)
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("单个输入文件不能超过 40 MB")
    return {"filename": path.name, "bytes": path.stat().st_size,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}


def _provenance():
    root = Path(__file__).parent
    return [{"filename": name, "sha256": hashlib.sha256((root / name).read_bytes()).hexdigest()}
            for name in ("bin_pixels.py", "solve_A.py", "web_bridge.py")]


def run_job(g_path, p_path, options=None):
    """Solve one selected G video against P and export accepted or rejected report.

    A has active target cells as rows, lag-major LC bin features as columns.
    The probability export restores all original target columns, including zeros.
    Rejected jobs expose only report.json and result.zip (which contains report).
    """
    start = time.perf_counter()
    options = dict(options or {})
    folder = Path(options.get("output_dir", "/tmp/lc_result"))
    folder.mkdir(parents=True, exist_ok=True)
    # Each invocation owns only these known files. Clear accepted artifacts from
    # a previous invocation before validation so rejection never serves stale A.
    names = ("A.csv", "A.npy", "A_factors.npz", "prediction.csv", "feature_metadata.csv",
             "target_cells.csv", "result.xlsx", "report.json", "result.zip")
    for name in names:
        (folder / name).unlink(missing_ok=True)
    summary = {"status": "rejected", "reasons": [], "metrics": {}, "input": {},
               "matrix_preview": [], "matrix_min": None, "matrix_max": None,
               "actual_preview": [], "prediction_preview": [], "frame_errors": [],
               "downloads": [], "elapsed_seconds": 0.0}
    report = {"schema_version": 1, "orientation": "A: (active P cells, LC features); prediction = F @ A.T",
              "formula": "Phi=F.T; Y=P.T; A0=Y@pinv(Phi); A=L@R with original numerical nonnegative guard",
              "scope": "Training frame interpolation only. New-frame nonnegativity, unit mass and generalization are not guaranteed.",
              "feature_order": "lag-major: all bins at lag 0, then lag 1, then lag 2; early history clamps to first frame",
              "source_provenance": _provenance(), "settings": options}
    files = []
    try:
        report["input_files"] = {"G": _source_info(g_path), "P": _source_info(p_path)}
        g_source, p_source = _load(g_path, "G", options), _load(p_path, "P", options)
        lc = _prepare_g(g_source, options)
        F = lc["F"][lc["selected_video_index"] - 1]
        p = p_source["P"]
        if p.ndim not in (2, 3) or min(p.shape) < 1 or len(p) > MAX_FRAMES:
            raise ValueError("P 必须为不超过 600 帧的 (T,K) 或 (T,H,W) 数组")
        cells = int(np.prod(p.shape[1:]))
        if "xy" in p_source and len(p_source["xy"]) != cells:
            raise ValueError("P 的 xy 坐标数量与网格数不一致")
        original_ids = p_source.get("active", np.arange(cells))
        p_source = dict(p_source, active=np.arange(cells))
        required_rank = options.get("required_rank")
        if required_rank not in (None, ""):
            required_rank = _number(required_rank, "required_rank")
            if required_rank < 1 or int(required_rank) != required_rank:
                raise ValueError("required_rank 必须为空或正整数")
            required_rank = int(required_rank)
        else:
            required_rank = None
        thresholds = {key: _number(options.get(key, 1e-10), key)
                      for key in ("max_fit_abs", "max_mass_error", "max_offset")}
        if any(value <= 0 for value in thresholds.values()):
            raise ValueError("验收阈值必须为正数")
        normalize = options.get("normalize_target", False)
        if not isinstance(normalize, (bool, np.bool_)):
            raise ValueError("normalize_target 必须为 true 或 false")
        target = prepare_target(p_source, lc["time"], lc["time_is_frame_index"],
                                options.get("time_align", "auto"), bool(normalize))
        active_indices = target["active"].astype(int)
        if len(active_indices) * F.shape[1] > MAX_ELEMENTS:
            raise ValueError("A 超过 2,000,000 个元素；请减少目标网格或 LC 分箱数量")
        features = _features(lc)
        xy = _target_xy(p_source)
        summary["input"] = {"frames": len(F), "bins": F.shape[1] // 3,
                            "features": F.shape[1], "cells": cells,
                            "active_cells": len(active_indices), "videos": len(lc["F"]),
                            "video_index": lc["selected_video_index"],
                            "video_id": str(lc["video_ids"][lc["selected_video_index"] - 1]),
                            "g_filename": Path(g_path).name, "p_filename": Path(p_path).name,
                            "g_shape": list(g_source.get("G", g_source.get("mean_green")).shape),
                            "p_shape": list(p.shape), "time_alignment": str(target["time_alignment"]),
                            "time_is_frame_index": bool(lc["time_is_frame_index"]),
                            "target_normalized": bool(normalize)}
        report.update(feature_metadata=features, active_indices=active_indices,
                      target_active_xy=xy[active_indices], source_target_ids=original_ids,
                      source_target_xy=xy, LC_time=lc["time"], model_time=target["model_time"],
                      target_cell_area=target["cell_area"],
                      original_target_mass_range=[float(target["original_mass_min"]), float(target["original_mass_max"])])
        F_plus, singular = input_pseudoinverse(F)
        arrays, prediction, metrics = solve_mapping(F, target, F_plus, singular)
        reasons = check_requirements(metrics, required_rank, **thresholds)
        summary["metrics"] = _plain(metrics)
        summary["reasons"] = reasons
        summary["frame_errors"] = [float(value) for value in np.max(np.abs(prediction - target["P"]), axis=1)]
        if not reasons:
            summary["status"] = "accepted"
            A = arrays["L"] @ arrays["R"]
            full_prediction = np.zeros((len(F), cells))
            full_actual = np.zeros_like(full_prediction)
            full_prediction[:, active_indices] = prediction
            full_actual[:, active_indices] = target["P"]
            summary.update(matrix_preview=A[:20, :20].tolist(), matrix_min=float(A.min()),
                           matrix_max=float(A.max()),
                           actual_preview=full_actual[0, :256].tolist(),
                           prediction_preview=full_prediction[0, :256].tolist(),
                           preview_xy=xy[:256].tolist(),
                           preview_cell_indices=list(range(min(cells, 256))),
                           preview_frame_index=0)
            np.save(folder / "A.npy", A, allow_pickle=False)
            np.savez_compressed(folder / "A_factors.npz", **arrays,
                                active_indices=active_indices, target_xy=xy[active_indices],
                                LC_time=lc["time"], model_time=target["model_time"],
                                feature_history_lag=np.array([r["history_lag"] for r in features]),
                                feature_bin_index=np.array([r["bin_index"] for r in features]))
            _write_csv(folder / "A.csv", ["target_cell_index"] + [f"feature_{i}" for i in range(A.shape[1])],
                       ([int(active_indices[i])] + list(row) for i, row in enumerate(A)))
            _write_csv(folder / "prediction.csv", ["LC_time"] + [f"p{i}" for i in range(cells)],
                       ([float(lc["time"][i])] + list(row) for i, row in enumerate(full_prediction)))
            feature_keys = ["feature_index", "history_lag", "bin_index", "kx", "ky", "pixel_count"]
            _write_csv(folder / "feature_metadata.csv", feature_keys,
                       ([record[key] for key in feature_keys] for record in features))
            active_set = set(map(int, active_indices))
            _write_csv(folder / "target_cells.csv", ["source_cell_index", "x", "y", "active", "A_row"],
                       ([i, float(coord[0]), float(coord[1]), int(i in active_set),
                         int(np.searchsorted(active_indices, i)) if i in active_set else ""] for i, coord in enumerate(xy)))
            files = [folder / name for name in names[:6]]
            if A.size + full_prediction.size * 2 + len(features) * 6 < MAX_XLSX_EXPORT:
                workbook = Workbook(write_only=True)
                sheet = workbook.create_sheet("A")
                sheet.append(["target_cell_index"] + [f"feature_{i}" for i in range(A.shape[1])])
                for i, row in enumerate(A):
                    sheet.append([int(active_indices[i])] + [float(v) for v in row])
                for title, values in (("prediction", full_prediction), ("actual_aligned", full_actual)):
                    sheet = workbook.create_sheet(title)
                    sheet.append(["LC_time"] + [f"p{i}" for i in range(cells)])
                    for i, row in enumerate(values):
                        sheet.append([float(lc["time"][i])] + [float(v) for v in row])
                sheet = workbook.create_sheet("features")
                sheet.append(feature_keys)
                for record in features:
                    sheet.append([record[key] for key in feature_keys])
                sheet = workbook.create_sheet("target_cells")
                sheet.append(["source_cell_index", "x", "y", "active"])
                for i, coord in enumerate(xy):
                    sheet.append([i, float(coord[0]), float(coord[1]), int(i in active_set)])
                workbook.save(folder / "result.xlsx")
                files.append(folder / "result.xlsx")
            else:
                report["xlsx_export_omitted"] = "超过 200,000 个单元格的导出上限；CSV/NPY/NPZ 包含完整数值"
    except (ValueError, TypeError, KeyError, IndexError, OSError, ArithmeticError,
            json.JSONDecodeError, zipfile.BadZipFile, np.linalg.LinAlgError) as error:
        summary["status"] = "rejected"
        summary["reasons"] = [str(error)]
        files = []
        for name in names[:7]:
            (folder / name).unlink(missing_ok=True)
    summary["elapsed_seconds"] = round(time.perf_counter() - start, 4)
    report.update(_plain({key: value for key, value in summary.items() if key != "downloads"}))
    _json(folder / "report.json", report)
    files.append(folder / "report.json")
    with zipfile.ZipFile(folder / "result.zip", "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in files:
            archive.write(path, path.name)
    files.append(folder / "result.zip")
    summary["downloads"] = [_download(path) for path in files]
    return _plain(summary)


def make_demo(directory="/tmp/lc_demo"):
    """A small reproducible full-row-rank G/P pair, with explicit physical times."""
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    rng = np.random.default_rng(20261016)
    frames, cells = 8, 12
    times = np.arange(frames) * 0.1
    xy = np.column_stack([np.arange(cells) * 16, np.zeros(cells)])
    green = rng.uniform(20, 230, (frames, cells))
    probability = rng.uniform(0.02, 1, (frames, 9))
    probability /= probability.sum(axis=1, keepdims=True)
    for role, values, coords in (("G", green, xy), ("P", probability,
                                 np.array([[x, y] for y in range(3) for x in range(3)]))):
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet("data")
        sheet.append(["t", "x", "y", role])
        for ti, row in zip(times, values):
            for coord, value in zip(coords, row):
                sheet.append([float(ti), float(coord[0]), float(coord[1]), float(value)])
        workbook.save(directory / f"demo_{role}.xlsx")
    return {"g_path": str(directory / "demo_G.xlsx"), "p_path": str(directory / "demo_P.xlsx")}


def make_templates(directory="/tmp/lc_templates"):
    """Return download descriptors for valid, editable XLSX, CSV and JSON examples."""
    directory = Path(directory)
    paths = make_demo(directory)
    files = [Path(paths["g_path"]), Path(paths["p_path"])]
    for role, key in (("G", "g_path"), ("P", "p_path")):
        rows, _ = _read_rows(Path(paths[key]))
        source = _table_source(rows, role)
        csv_path = directory / f"template_{role}.csv"
        _write_csv(csv_path, rows[0], rows[1:])
        json_path = directory / f"template_{role}.json"
        _json(json_path, source)
        files.extend([csv_path, json_path])
    return [_download(path) for path in files]
