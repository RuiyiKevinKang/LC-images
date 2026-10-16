"""Integration checks for the upload bridge, using the original solver as oracle."""
import csv
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1] / "site/python"
sys.path.insert(0, str(ROOT))
import web_bridge as bridge
from bin_pixels import prepare_lc
from solve_A import prepare_target, input_pseudoinverse, solve_mapping


class BridgeIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.folder = Path(self.temporary.name)

    def tearDown(self):
        assert self.folder.resolve().is_relative_to(Path(tempfile.gettempdir()).resolve())
        self.temporary.cleanup()

    def write_json(self, name, data):
        path = self.folder / name
        path.write_text(json.dumps(bridge._plain(data)), encoding="utf-8")
        return path

    def options(self, **extra):
        return {"output_dir": str(self.folder / "result"), **extra}

    def pair(self, frames=5):
        rng = np.random.default_rng(42)
        G = rng.uniform(1, 200, (frames, 8))
        P = rng.uniform(0.02, 1, (frames, 6))
        P /= P.sum(axis=1, keepdims=True)
        xy = np.column_stack([np.arange(8) * 16, np.zeros(8)])
        time = np.arange(frames) * .1
        gp = self.write_json("G.json", {"G": G, "xy": xy, "time": time})
        pp = self.write_json("P.json", {"P": P, "time": time})
        return gp, pp, G, P, xy, time

    def test_numerical_parity_and_restored_zero_cells(self):
        gp, pp, G, P, xy, t = self.pair()
        P = np.column_stack([P[:, :2], np.zeros(len(P)), P[:, 2:]])
        pp = self.write_json("P.json", {"P": P, "time": t})
        result = bridge.run_job(gp, pp, self.options())
        self.assertEqual(result["status"], "accepted", result["reasons"])
        lc = prepare_lc(G, xy, t)
        target = prepare_target({"P": P, "time": t}, t)
        arrays, prediction, metrics = solve_mapping(lc["F"][0], target,
                                                    *input_pseudoinverse(lc["F"][0]))
        matrix = np.load(self.folder / "result/A.npy")
        np.testing.assert_array_equal(matrix, arrays["L"] @ arrays["R"])
        self.assertEqual(result["metrics"], metrics)
        self.assertEqual(result["input"]["cells"], 7)
        self.assertEqual(result["input"]["active_cells"], 6)
        with (self.folder / "result/prediction.csv").open(encoding="utf-8-sig") as stream:
            rows = list(csv.reader(stream))
        self.assertEqual(rows[0][3], "p2")
        self.assertTrue(all(float(row[3]) == 0 for row in rows[1:]))
        report = json.loads((self.folder / "result/report.json").read_text(encoding="utf-8"))
        self.assertEqual(report["active_indices"], [0, 1, 3, 4, 5, 6])
        self.assertEqual(len(report["source_provenance"]), 3)
        json.dumps(result, allow_nan=False)

    def test_xlsx_demo_roundtrip_and_templates(self):
        demo = bridge.make_demo(self.folder / "demo")
        result = bridge.run_job(demo["g_path"], demo["p_path"], self.options())
        self.assertEqual(result["status"], "accepted", result["reasons"])
        self.assertEqual(bridge.inspect_input(demo["g_path"], "G")["frames"], 8)
        workbook = load_workbook(self.folder / "result/result.xlsx", read_only=True, data_only=True)
        sheet = workbook["A"]
        matrix = np.array([row[1:] for row in sheet.iter_rows(min_row=2, values_only=True)])
        np.testing.assert_allclose(matrix, np.load(self.folder / "result/A.npy"), rtol=1e-14, atol=1e-14)
        workbook.close()
        descriptors = bridge.make_templates(self.folder / "templates")
        self.assertEqual(len(descriptors), 6)
        for item in descriptors:
            role = "G" if "_G." in item["name"] else "P"
            self.assertEqual(bridge.inspect_input(item["path"], role)["frames"], 8)

    def test_rejected_rank_removes_old_outputs(self):
        gp, pp, G, P, xy, t = self.pair()
        self.assertEqual(bridge.run_job(gp, pp, self.options())["status"], "accepted")
        gp = self.write_json("G.json", {"G": np.ones_like(G), "xy": xy, "time": t})
        result = bridge.run_job(gp, pp, self.options())
        self.assertEqual(result["status"], "rejected")
        self.assertIn("F秩=", result["reasons"][0])
        self.assertFalse((self.folder / "result/A.npy").exists())
        self.assertEqual([d["name"] for d in result["downloads"]], ["report.json", "result.zip"])
        with zipfile.ZipFile(self.folder / "result/result.zip") as archive:
            self.assertEqual(archive.namelist(), ["report.json"])

    def test_long_table_duplicate_and_missing_coordinates(self):
        path = self.folder / "bad.csv"
        for content in ("t,x,y,G\n0,0,0,1\n0,0,0,2\n1,0,0,3\n1,1,0,4\n",
                        "t,x,y,G\n0,0,0,1\n0,1,0,2\n1,0,0,3\n"):
            path.write_text(content, encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "坐标|重复"):
                bridge.inspect_input(path, "G")

    def test_time_and_probability_semantics(self):
        gp, pp, G, P, xy, t = self.pair()
        pp = self.write_json("P.json", {"P": P})
        result = bridge.run_job(gp, pp, self.options())
        self.assertEqual(result["status"], "rejected")
        self.assertIn("只有一侧提供时间", result["reasons"][0])
        self.assertEqual(bridge.run_job(gp, pp, self.options(time_align="frame"))["status"], "accepted")
        pp = self.write_json("P.json", {"P": P * 2, "time": t})
        self.assertEqual(bridge.run_job(gp, pp, self.options())["status"], "rejected")
        self.assertEqual(bridge.run_job(gp, pp, self.options(normalize_target=True))["status"], "accepted")
        self.assertEqual(bridge.run_job(gp, pp, self.options(target_kind="density"))["status"], "rejected")
        self.assertEqual(bridge.run_job(gp, pp, self.options(target_kind="density", cell_area=.5))["status"], "accepted")

    def test_npz_multiple_video_and_cached_binning(self):
        gp, pp, G, P, xy, t = self.pair()
        lc = prepare_lc(G, xy, t)
        path = self.folder / "cache.npz"
        np.savez(path, mean_green=np.concatenate([lc["mean_green"], lc["mean_green"] * 2]),
                 time=t, video_ids=np.array(["first", "second"]))
        info = bridge.inspect_input(path, "G", {"video_index": 2})
        self.assertEqual(info["videos"], 2)
        self.assertTrue(info["is_binned_cache"])
        result = bridge.run_job(path, pp, self.options(video_index=2))
        self.assertEqual(result["status"], "accepted", result["reasons"])
        self.assertEqual(result["input"]["video_id"], "second")
        self.assertEqual(bridge.run_job(path, pp, self.options(video_index=3))["status"], "rejected")

    def test_nonfinite_complex_and_object_rejected(self):
        gp, pp, G, P, xy, t = self.pair()
        for array in (np.array([[[complex(1, 1)]]]), np.array([[[np.inf]]]), np.array([[[1]]], dtype=object)):
            path = self.folder / "bad.npy"
            np.save(path, array)
            self.assertEqual(bridge.run_job(path, pp, self.options())["status"], "rejected")

    def test_linear_time_alignment_and_no_extrapolation(self):
        gp, pp, G, P, xy, t = self.pair()
        target_times = np.array([t[0], t[-1]])
        pp = self.write_json("P.json", {"P": P[[0, -1]], "time": target_times})
        result = bridge.run_job(gp, pp, self.options(time_align="linear"))
        self.assertEqual(result["status"], "accepted", result["reasons"])
        report = json.loads((self.folder / "result/report.json").read_text(encoding="utf-8"))
        self.assertEqual(report["input"]["time_alignment"], "linear")
        pp = self.write_json("P.json", {"P": P[[0, -1]], "time": [t[1], t[-1]]})
        result = bridge.run_job(gp, pp, self.options(time_align="linear"))
        self.assertEqual(result["status"], "rejected")
        self.assertIn("不允许外推", result["reasons"][0])

    def test_wide_tables_and_required_rank_rejection(self):
        gp, pp, G, P, xy, t = self.pair()
        gp = self.folder / "wide_G.tsv"
        with gp.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.writer(stream, delimiter="\t")
            writer.writerow(["pixel_id", "x", "y"] + [f"G_frame_{i:03d}" for i in range(len(G))])
            for i in range(G.shape[1]):
                writer.writerow([i, *xy[i], *G[:, i]])
        pp = self.folder / "wide_P.csv"
        with pp.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow([f"p{i}" for i in range(P.shape[1])])
            writer.writerows(P)
        result = bridge.run_job(gp, pp, self.options())
        self.assertEqual(result["status"], "accepted", result["reasons"])
        self.assertEqual(result["input"]["time_alignment"], "frame")
        result = bridge.run_job(gp, pp, self.options(required_rank=1))
        self.assertEqual(result["status"], "rejected")
        self.assertIn("不满足要求", result["reasons"][0])
        self.assertEqual(len(result["downloads"]), 2)

    def test_oversized_npy_header_rejected_before_allocation(self):
        gp, pp, G, P, xy, t = self.pair()
        path = self.folder / "oversized.npy"
        with path.open("wb") as stream:
            np.lib.format.write_array_header_1_0(stream, {
                "shape": (10**12,), "fortran_order": False, "descr": "<f8"})
        result = bridge.run_job(path, pp, self.options())
        self.assertEqual(result["status"], "rejected")
        self.assertIn("元素数量", result["reasons"][0])

    def test_npy_images_and_npz_boolean_roi_mask(self):
        rng = np.random.default_rng(52)
        G = rng.uniform(0, 255, (5, 33, 33))
        P = rng.uniform(.01, 1, (5, 3, 3))
        P /= P.sum(axis=(1, 2), keepdims=True)
        gp, pp = self.folder / "G.npy", self.folder / "P.npy"
        np.save(gp, G)
        np.save(pp, P)
        result = bridge.run_job(gp, pp, self.options())
        self.assertEqual(result["status"], "accepted", result["reasons"])
        self.assertEqual(result["input"]["time_alignment"], "frame")
        mask = np.ones((33, 33), dtype=bool)
        mask[:8] = False
        gp = self.folder / "masked.npz"
        np.savez(gp, G=G, mask=mask)
        result = bridge.run_job(gp, pp, self.options())
        self.assertEqual(result["status"], "accepted", result["reasons"])
        self.assertEqual(result["input"]["cells"], 9)


if __name__ == "__main__":
    unittest.main(verbosity=2)
